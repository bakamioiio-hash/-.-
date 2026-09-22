import os
import re
import time
import urllib.parse
from html import unescape
import requests
from openpyxl import Workbook, load_workbook

URL_FILE = "data/urls.txt"
EXCEL_FILE = "data/leads.xlsx"

# 劳保产品专属关键词池（覆盖安全鞋、防护手套、反光服、工业安全防护，按日轮换）
KEYWORDS_POOL = [
    # --- 综合 PPE 与工业安全品分销商 ---
    "safety supplies distributor USA wholesale inquiry",
    "PPE personal protective equipment dealer UK contact us",
    "industrial safety equipment distributor Germany contact email",
    "workplace safety gear wholesaler Australia contact",
    "safety products importer distributor Canada email",
    "occupational safety equipment supplier Europe contact us",
    
    # --- 劳保手套（防割、丁腈、浸胶、耐磨手套）---
    "protective work gloves distributor USA wholesale email",
    "cut resistant gloves supplier UK contact us",
    "industrial safety gloves wholesaler Germany inquiry",
    "heavy duty work gloves importer Canada contact",
    "disposable nitrile gloves wholesale distributor Europe email",
    "leather work gloves supplier Australia contact us",
    
    # --- 劳保工装与反光服（Hi-Vis / Flame Resistant / Coveralls）---
    "hi vis workwear safety clothing distributor USA contact us",
    "industrial workwear and uniforms supplier UK email",
    "flame resistant FR clothing distributor Canada inquiry",
    "high visibility reflective vest supplier Australia contact",
    "construction workwear apparel dealer Germany contact",
    "protective coveralls safety workwear Europe email",
    
    # --- 劳保鞋与安全靴（Safety Shoes / Steel Toe Boots）---
    "safety shoes work boots distributor USA wholesale contact",
    "steel toe work boots supplier UK inquiry email",
    "safety footwear importer distributor Australia contact us",
    "industrial safety shoes dealer Germany contact email",
    "protective work boots wholesaler Canada contact",
    "occupational safety footwear brand Europe email",
    
    # --- 头部/眼部/呼吸防护（Helmets / Goggles / Masks）---
    "safety glasses eye protection supplier USA contact us",
    "hard hat safety helmet distributor UK inquiry",
    "welding protective gear and masks distributor Europe email",
    "hearing and eye protection equipment supplier Australia contact",
    
    # --- 细分特种劳保 ---
    "waterproof industrial rainwear supplier UK contact us",
    "fall protection safety harness distributor USA email",
    "ESD antistatic cleanroom clothing supplier Europe contact",
    "road work safety apparel distributor Canada contact us"
]

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
}

def init_excel():
    os.makedirs("data", exist_ok=True)
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Leads"
        ws.append(["抓取时间", "目标独立站", "商业邮箱", "电话/WhatsApp", "社媒主页", "品类与网站简介"])
        wb.save(EXCEL_FILE)

def get_existing_urls():
    if not os.path.exists(EXCEL_FILE):
        return set()
    try:
        wb = load_workbook(EXCEL_FILE, read_only=True)
        ws = wb.active
        urls = {row[1] for row in ws.iter_rows(values_only=True) if row and row[1]}
        return urls
    except Exception:
        return set()

def fetch_search_urls(query, max_results=10):
    urls = []
    try:
        search_url = f"https://html.duckduckgo.com/html/?q={urllib.parse.quote(query)}"
        resp = requests.get(search_url, headers=HEADERS, timeout=15)
        matches = re.findall(r'<a class="result__url"[^>]*href="([^"]+)"', resp.text)
        for m in matches:
            if "duckduckgo.com/l/?uddg=" in m:
                target = urllib.parse.unquote(m.split("uddg=")[1].split("&")[0])
            else:
                target = m.strip()
            if target.startswith("http") and not any(x in target for x in ["amazon.", "alibaba.", "aliexpress.", "ebay."]):
                base_domain = urllib.parse.urlparse(target).scheme + "://" + urllib.parse.urlparse(target).netloc
                if base_domain not in urls:
                    urls.append(base_domain)
            if len(urls) >= max_results:
                break
    except Exception as e:
        print(f"搜索获取链接异常: {e}")
    return urls

def extract_page_info(url):
    try:
        resp = requests.get(url, headers=HEADERS, timeout=12)
        text = resp.text
        
        # 邮箱匹配并过滤静态资源后缀
        raw_emails = re.findall(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', text)
        clean_emails = [e for e in set(raw_emails) if not any(e.lower().endswith(x) for x in ['.png', '.jpg', '.jpeg', '.gif', '.svg', '.webp'])]
        
        # 电话 / WhatsApp
        phones = re.findall(r'(\+?[0-9]{1,3}[-.\s]?\(?[0-9]{1,4}\)?[-.\s]?[0-9]{3,4}[-.\s]?[0-9]{3,4})', text)
        clean_phones = list(set([p.strip() for p in phones if len(re.sub(r'\D', '', p)) >= 8]))[:2]

        # 社媒
        socials = []
        for plat in ["instagram.com", "facebook.com", "linkedin.com", "tiktok.com"]:
            matches = re.findall(rf'(https?://(?:www\.)?{plat}/[a-zA-Z0-9_\-\.]+)', text)
            if matches:
                socials.append(matches[0])

        # 简介提取
        desc = ""
        meta_desc = re.findall(r'<meta[^>]*name=["\']description["\'][^>]*content=["\'](.*?)["\']', text, re.I)
        if meta_desc:
            desc = unescape(meta_desc[0].strip())
        elif "<title>" in text:
            desc = unescape(re.findall(r'<title>(.*?)</title>', text, re.I)[0].strip())

        return ", ".join(clean_emails), ", ".join(clean_phones), ", ".join(set(socials)), desc[:200]
    except Exception:
        return "", "", "", ""

def run():
    init_excel()
    existing_urls = get_existing_urls()
    
    # 轮换关键词
    keyword_index = int(time.time() / 86400) % len(KEYWORDS_POOL)
    current_kw = KEYWORDS_POOL[keyword_index]
    print(f"[*] 今日检索关键词: {current_kw}")

    target_urls = []
    if os.path.exists(URL_FILE):
        with open(URL_FILE, "r", encoding="utf-8") as f:
            target_urls = [line.strip() for line in f if line.strip()]

    if not target_urls:
        target_urls = fetch_search_urls(current_kw, max_results=10)

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    added = 0

    for idx, url in enumerate(target_urls, start=1):
        if url in existing_urls:
            continue
        print(f"[{idx}/{len(target_urls)}] 抓取: {url}")
        emails, phones, socials, desc = extract_page_info(url)
        
        ws.append([
            time.strftime("%Y-%m-%d %H:%M"),
            url,
            emails or "未公开",
            phones or "未公开",
            socials or "未公开",
            desc or "暂无主营简介"
        ])
        wb.save(EXCEL_FILE)
        existing_urls.add(url)
        added += 1
        print(f"  + 邮箱: {emails or '无'} | 社媒: {socials or '无'}")
        time.sleep(2)

    print(f"[✓ 任务完成! 本次新增有效线索 {added} 条，已追加至 {EXCEL_FILE}]")

if __name__ == "__main__":
    run()