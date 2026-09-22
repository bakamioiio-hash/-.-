import os
import subprocess
import sys
from openpyxl import load_workbook

def run_step(desc, cmd):
    print(f"\n{'='*20} {desc} {'='*20}")
    print(f"[CMD] {cmd}")
    res = subprocess.run(cmd, shell=True)
    if res.returncode != 0:
        print(f"[!] 执行异常，退出码: {res.returncode}")
        sys.exit(res.returncode)

def main():
    # 步骤 1：执行海关反查入库
    run_step("步骤 1：执行海关反查与提单买家解析", "python tools/customs_leads_extractor.py --mode auto")

    # 步骤 2：检查 Excel 落盘情况
    excel_path = "data/leads.xlsx"
    print(f"\n{'='*20} 步骤 2：校验 data/leads.xlsx 数据底册 {'='*20}")
    if not os.path.exists(excel_path):
        print("[!] 错误：未找到 data/leads.xlsx")
        return

    wb = load_workbook(excel_path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    print(f"[✓] 成功读取 Excel，当前总行数（含表头）: {len(rows)}")
    print("最新 2 条录入买家数据预览：")
    for r in rows[-2:]:
        print(f"  - 品类: {r[1]} | 国家: {r[2]} | 买家: {r[3]} | 邮箱: {r[5]} | 状态: {r[8]}")

    # 步骤 3：测试开发信装配引擎
    run_step("步骤 3：加载底册并预览针对性外贸开发信", "python tools/send_cold_emails.py --mode preview")

    print("\n[✓] 全链路测试完成！各模块运行正常。")

if __name__ == "__main__":
    main()